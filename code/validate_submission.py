"""Check consistency of code, spreadsheet, plots and official evaluation artifacts."""
import ast
import json
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from results_table import load_results

def validate_submission(out,repo):
    out,repo=Path(out),Path(repo)
    required=("REPORT.md","experiments.xlsx","predictions_eval.csv","eval_result.json",
              "figures","code/lab.ipynb","run_manifest.json")
    for name in required: assert (out/name).exists(),f"Missing {name}"
    for name in ("data.py","model.py","optimizer.py","train.py","plots.py","results_table.py"):
        assert (out/"code"/name).exists(),name
    for p in (out/"code").glob("*.py"):
        tree=ast.parse(p.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node,ast.Raise) and node.exc is not None:
                assert not (isinstance(node.exc,ast.Name) and node.exc.id=="NotImplementedError"),str(p)
    results=load_results(out/"results")
    wb=openpyxl.load_workbook(out/"experiments.xlsx",data_only=False)
    cached=openpyxl.load_workbook(out/"experiments.xlsx",data_only=True)
    assert wb.sheetnames==["Legend","Experiments","Seeds","Summary"]
    headers=[c.value for c in wb["Experiments"][1]]
    rows=[dict(zip(headers,row)) for row in cached["Experiments"].iter_rows(min_row=2,values_only=True) if row[0]]
    ids=[r["exp_id"] for r in rows]
    assert len(ids)==len(set(ids))==len(results)
    assert set(ids)=={r["cfg"]["exp_id"] for r in results}
    indexed={r["exp_id"]:r for r in rows}
    for result in results:
        eid=result["cfg"]["exp_id"]
        assert (out/"figures"/f"{eid}.png").is_file(),eid
        assert (out/"figures"/f"{eid}.png").stat().st_size>1000
        row=indexed[eid]
        for key in ("val_acc","val_macro_f1","step0_loss","best_val_loss"):
            expected=result["summary"].get(key)
            if expected is not None: assert abs(row[key]-expected)<1e-10,(eid,key)
        assert row["figure_file"]==f"figures/{eid}.png"
        assert len(result["history"]["epoch"])==result["summary"]["completed_epochs"]
    for sheet in cached:
        for line in sheet:
            for cell in line: assert cell.data_type!="e",(sheet.title,cell.coordinate,cell.value)
    formula_cols=("step0_gap_vs_lnC","gap_val_minus_train","delta_val_f1_vs_base","beyond_noise")
    for i in range(2,len(rows)+2):
        for key in formula_cols: assert wb["Experiments"].cell(i,headers.index(key)+1).data_type=="f"
    manifest=json.loads((out/"run_manifest.json").read_text(encoding="utf-8"))
    assert manifest["selection_fixed_before_eval"]
    pred=pd.read_csv(out/"predictions_eval.csv")
    assert pred.columns.tolist()==["row_id","pred"]
    with np.load(repo/"data/processed/eval.npz") as data:
        assert len(pred)==len(data["row_id"])==116203
        assert pred.row_id.is_unique and set(pred.row_id)==set(data["row_id"])
        aligned=pred.set_index("row_id").loc[data["row_id"],"pred"].to_numpy()
        assert np.array_equal(aligned,aligned.astype(int)) and ((aligned>=0)&(aligned<7)).all()
        cm=np.zeros((7,7),dtype=np.int64)
        np.add.at(cm,(data["y"],aligned.astype(int)),1)
    scores=json.loads((out/"eval_result.json").read_text())
    assert np.array_equal(cm,np.asarray(scores["confusion_matrix"]))
    assert abs(np.trace(cm)/cm.sum()-scores["accuracy"])<1e-12
    selected=indexed[manifest["selected_exp_id"]]
    assert abs(selected["eval_macro_f1"]-scores["macro_f1"])<1e-12
    assert abs(selected["eval_acc"]-scores["accuracy"])<1e-12
    for row in rows:
        if row["exp_id"] not in manifest.get("eval_files",{}):
            assert row["eval_macro_f1"] is None and row["eval_acc"] is None
        else:
            official=json.loads((out/manifest["eval_files"][row["exp_id"]]).read_text())
            assert abs(row["eval_macro_f1"]-official["macro_f1"])<1e-12
    assert not list(out.rglob("*.pt")) and not list(out.rglob("*.npz"))
    print(f"Submission validation passed: {len(rows)} experiment rows/plots, four-sheet formulas and caches, 116203 predictions, official scores matched.")
    return True

if __name__=="__main__":
    import sys
    from workflow import find_repo
    validate_submission(sys.argv[1],find_repo())
